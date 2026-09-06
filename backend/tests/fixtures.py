"""Programmatic test fixtures: build each format's bytes at test time.

Avoids committing binaries and keeps the "golden" expectations in code next to
the assertions.
"""

from __future__ import annotations

import io

_FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def make_markdown(title: str = "Intro", body: str = "hello markdown world") -> bytes:
    return f"# {title}\n\n{body}\n".encode()


def make_text(body: str = "plain text line one\nplain text line two") -> bytes:
    return body.encode("utf-8")


def make_pdf(text: str = "hello pdf world") -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=18)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(heading: str = "Section One", body: str = "hello docx world") -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "k"
    table.rows[0].cells[1].text = "v"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_scanned_png(text: str = "SCANNED DOCUMENT") -> bytes:
    """Render text onto a white image so Tesseract can OCR it back."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (900, 200), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype(_FONT_PATH, 64)
    except OSError:  # pragma: no cover - font always present in CI/container
        font = ImageFont.load_default()
    draw.text((40, 60), text, fill="black", font=font)
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()
