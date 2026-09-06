"""Lightweight DOCX parser (python-docx).

In production, Office formats route through Docling (layout + tables). This
dependency-light adapter is the fast path and the fallback when Docling is
unavailable, and keeps DOCX ingestion fully local and testable. Paragraph
styles named ``Heading N`` are preserved as ``heading_path``.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from app.domain.models import ParsedDocument, TableBlock, TextBlock

if TYPE_CHECKING:
    from app.core.config import Settings

_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


class DocxParser:
    """Extracts paragraphs and tables from a .docx via python-docx."""

    name = "docx"

    @classmethod
    def from_settings(cls, settings: Settings) -> DocxParser:
        return cls()

    def supports(self, mime_type: str) -> bool:
        return mime_type in _MIME_TYPES

    async def parse(
        self, data: bytes, *, filename: str, mime_type: str
    ) -> ParsedDocument:
        import docx  # lazy import

        document = docx.Document(io.BytesIO(data))

        blocks: list[TextBlock] = []
        heading_path: str | None = None
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style else ""
            if style.lower().startswith("heading"):
                heading_path = text
            blocks.append(TextBlock(text=text, page=1, heading_path=heading_path))

        tables: list[TableBlock] = []
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                tables.append(TableBlock(rows=rows, page=1))

        return ParsedDocument(
            source_filename=filename,
            mime_type=mime_type,
            text_blocks=blocks,
            tables=tables,
            page_count=1,
            parser_name=self.name,
        )
