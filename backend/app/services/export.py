"""Export an article to PDF, DOCX or Markdown (returns raw bytes)."""
from __future__ import annotations

import io
import re

from app.models.article import Article


def _source_label(src: dict) -> str:
    title = src.get("title") or "Untitled source"
    url = src.get("url")
    stype = src.get("type")
    label = f"{title} ({url})" if url else title
    return f"[{stype}] {label}" if stype else label


def to_markdown(article: Article) -> bytes:
    kw = ", ".join(article.keywords or [])
    tags = ", ".join(article.tags or [])
    front = (
        f"# {article.title}\n\n"
        f"> {article.summary}\n\n"
        f"**Keywords:** {kw}\n\n"
        f"**Tags:** {tags}\n\n"
        f"**Sentiment:** {article.sentiment}\n\n"
        "---\n\n"
    )
    body = article.body or ""
    if article.sources:
        lines = "\n".join(
            f"- {_source_label(s)}" for s in article.sources if isinstance(s, dict)
        )
        body += f"\n\n## Sources\n\n{lines}\n"
    return (front + body).encode("utf-8")


def _strip_md(text: str) -> str:
    """Very light Markdown -> plain text for PDF/DOCX bodies."""
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`{1,3}(.*?)`{1,3}", r"\1", text)
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", text)
    return text


def to_docx(article: Article) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(article.title or "Untitled", level=0)
    if article.summary:
        p = doc.add_paragraph(article.summary)
        p.runs[0].italic = True

    for block in (article.body or "").split("\n\n"):
        block = block.strip()
        if not block:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", block)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(heading.group(2), level=level)
        else:
            doc.add_paragraph(_strip_md(block))

    if article.keywords:
        doc.add_heading("Keywords", level=2)
        doc.add_paragraph(", ".join(article.keywords)).runs[0].font.size = Pt(10)

    if article.sources:
        doc.add_heading("Sources", level=2)
        for s in article.sources:
            if isinstance(s, dict):
                doc.add_paragraph(_source_label(s), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(article: Article) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, title=article.title or "Untitled")
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=11, leading=16, alignment=TA_LEFT
    )

    flow = [Paragraph(_escape(article.title or "Untitled"), styles["Title"])]
    if article.summary:
        flow.append(Paragraph(f"<i>{_escape(article.summary)}</i>", styles["Italic"]))
    flow.append(Spacer(1, 12))

    for block in (article.body or "").split("\n\n"):
        block = block.strip()
        if not block:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", block)
        if heading:
            level = min(len(heading.group(1)), 4)
            flow.append(Paragraph(_escape(heading.group(2)), styles[f"Heading{level}"]))
        else:
            flow.append(Paragraph(_escape(_strip_md(block)), body_style))
        flow.append(Spacer(1, 6))

    if article.sources:
        flow.append(Spacer(1, 12))
        flow.append(Paragraph("Sources", styles["Heading2"]))
        for s in article.sources:
            if isinstance(s, dict):
                flow.append(Paragraph(f"• {_escape(_source_label(s))}", body_style))

    doc.build(flow)
    return buf.getvalue()


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


EXPORTERS = {
    "md": (to_markdown, "text/markdown", "md"),
    "markdown": (to_markdown, "text/markdown", "md"),
    "docx": (
        to_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "pdf": (to_pdf, "application/pdf", "pdf"),
}
