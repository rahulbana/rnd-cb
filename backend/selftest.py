"""Offline smoke test: parsing + HTML rendering, no OpenAI call."""
import io

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.document_parser import extract_text
from app.html_generator import render_html
from app.schemas import (
    StudyMaterial, Notes, NoteSection, Questions, TrueFalseQ, MCQ,
    FillBlankQ, ShortLongQ, CaseBasedQ, CaseQuestion,
)
from app.main import app  # ensure the FastAPI app imports cleanly


def test_text():
    assert "Photosynthesis" in extract_text("a.txt", b"Photosynthesis is a process.")


def test_docx():
    doc = Document()
    doc.add_paragraph("The water cycle has evaporation and condensation.")
    buf = io.BytesIO()
    doc.save(buf)
    assert "water cycle" in extract_text("a.docx", buf.getvalue())


def test_pptx():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Newton's laws of motion"
    buf = io.BytesIO()
    prs.save(buf)
    assert "Newton" in extract_text("a.pptx", buf.getvalue())


def test_html_render():
    material = StudyMaterial(
        notes=Notes(
            title="The Water Cycle",
            summary="An overview of how water moves.",
            key_points=["Evaporation", "Condensation", "Precipitation"],
            sections=[NoteSection(heading="Stages", points=["Rain", "Runoff"])],
            glossary=["Evaporation: liquid to vapour"],
        ),
        questions=Questions(
            true_false=[TrueFalseQ(statement="Rain is precipitation.", answer=True, explanation="Yes.")],
            mcq=[MCQ(question="Which is a gas?", options=["Ice", "Vapour", "Water"], answer="Vapour", explanation="Vapour is gaseous.")],
            fill_blanks=[FillBlankQ(question="Water ______ into vapour.", answer="evaporates")],
            very_short=[ShortLongQ(question="Define condensation.", answer="Vapour to liquid.")],
            short=[ShortLongQ(question="Explain runoff.", answer="Water flowing over land.")],
            long=[ShortLongQ(question="Describe the water cycle.", answer="A long answer...")],
            case_based=[CaseBasedQ(case="A pond dries in summer.", questions=[CaseQuestion(question="Why?", answer="Evaporation.")])],
        ),
    )
    html = render_html(material, "science.pdf", "Class 6")
    for token in ["The Water Cycle", "True", "Vapour", "evaporates", "Show answer", "<!DOCTYPE html>"]:
        assert token in html, f"missing {token!r}"
    # Escaping check
    hostile = render_html(
        StudyMaterial(notes=Notes(title="<script>x</script>")), "f", ""
    )
    assert "<script>x</script>" not in hostile
    assert "&lt;script&gt;" in hostile


def test_variants():
    from app.html_generator import render_variants
    material = StudyMaterial(
        notes=Notes(title="Cells", summary="About cells.", key_points=["Unit of life"]),
        questions=Questions(
            mcq=[MCQ(question="Powerhouse?", options=["Nucleus", "Mitochondria"],
                     answer="Mitochondria", explanation="ATP")],
            short=[ShortLongQ(question="Define cell.", answer="Basic unit of life.")],
        ),
    )
    v = render_variants(material, "bio.pdf", "Class 8")
    assert set(v) == {"notes", "questions_with_answers", "questions_only"}

    # Notes doc: has notes, no questions/answers
    assert "About cells." in v["notes"]
    assert "Powerhouse?" not in v["notes"]

    # Q + A: has questions AND the answer key
    assert "Powerhouse?" in v["questions_with_answers"]
    assert "Mitochondria" in v["questions_with_answers"]
    assert "Basic unit of life." in v["questions_with_answers"]

    # Questions only: has questions but NOT the answers
    qo = v["questions_only"]
    assert "Powerhouse?" in qo
    assert "Basic unit of life." not in qo   # answer text absent
    assert "Show answer" not in qo           # no answer blocks
    assert "Nucleus" in qo and "Mitochondria" in qo  # MCQ options kept


if __name__ == "__main__":
    test_text(); print("text OK")
    test_docx(); print("docx OK")
    test_pptx(); print("pptx OK")
    test_html_render(); print("html OK")
    test_variants(); print("variants OK")
    print("ALL PASSED")
